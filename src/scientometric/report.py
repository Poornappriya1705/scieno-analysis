from pathlib import Path
import math
import pandas as pd
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def fmt(v):
    if pd.isna(v): return "-"
    if isinstance(v,float): return f"{v:,.2f}"
    if isinstance(v,int): return f"{v:,}"
    return str(v)

def shade(cell,fill="D9EAF7"):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def add_table(doc,df,caption,max_rows=30):
    doc.add_paragraph(caption,style="Caption")
    d=df.head(max_rows).copy()
    t=doc.add_table(rows=1,cols=len(d.columns)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,c in enumerate(d.columns):
        t.rows[0].cells[i].text=str(c).replace("_"," "); shade(t.rows[0].cells[i])
    for _,r in d.iterrows():
        cells=t.add_row().cells
        for i,v in enumerate(r): cells[i].text=fmt(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return t

def add_figure(doc,path,caption):
    if Path(path).exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(6.4))
        cp=doc.add_paragraph(caption,style="Caption"); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER

def interpretation(name,df=None,overview=None):
    if name=="overview":
        o=overview; return f"The dataset contains {o['documents']:,} unique publications spanning {o['timespan']}. These records accumulated {o['citations']:,} citations, averaging {o['mean_citations']:.2f} citations per publication, with a collection-level h-index of {o['h_index']}. The annual growth rate ({o['annual_growth_rate']:.2f}%) is the compound annual growth rate (CAGR) calculated from the first and last annual publication counts across the observed timespan."
    if df is None or df.empty:return "The available metadata were insufficient to calculate this indicator reliably."
    top=df.iloc[0]
    label=df.columns[0]; val=top.iloc[0]
    if "Documents" in df.columns:return f"{val} ranks first by productivity with {int(top['Documents']):,} publications. Its citation total is {int(top.get('Citations',0)):,}, while mean citations distinguish productivity from average scholarly influence. Rankings should therefore be interpreted jointly rather than treating publication count alone as research quality."
    if "Frequency" in df.columns:return f"The most frequent item is '{val}' with {int(top['Frequency']):,} occurrences. High frequency indicates sustained attention within the corpus; it does not by itself establish novelty or causal importance. Co-occurrence and temporal analyses provide stronger evidence about thematic structure and evolution."
    return "The table summarizes the distribution and relative prominence of the measured entities in the dataset."

def build_report(cfg,results,figures,outfile):
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.8); sec.bottom_margin=Inches(.8); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
    styles=doc.styles; styles["Normal"].font.name="Times New Roman"; styles["Normal"].font.size=Pt(12)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(cfg["project"]["chapter_title"]); r.bold=True; r.font.size=Pt(16)
    doc.add_heading("4.1 Introduction",level=1)
    doc.add_paragraph("This chapter presents a reproducible scientometric analysis of bibliographic records retrieved from Web of Science and Scopus. The analysis covers publication growth, citation impact, productive authors and sources, country and institutional participation, authorship patterns, keyword structure, collaboration networks, bibliometric laws, and machine-assisted topic discovery. Each indicator is followed by a data-based explanation and a methodological caution to support defensible interpretation.")
    o=results["overview"]
    doc.add_heading("4.2 Dataset Profile and Data Preparation",level=1)
    prep=pd.DataFrame({"Metric":["Raw records","Unique cleaned records","Duplicates removed","Timespan","Databases"],"Value":[results['cleaning']['raw_records'],results['cleaning']['clean_records'],results['cleaning']['duplicates_removed'],o['timespan'],", ".join(sorted(results['databases']))]})
    add_table(doc,prep,"Table 4.1 Dataset preparation summary")
    doc.add_paragraph("Duplicate removal prioritised DOI and then normalized title. Database provenance was retained because citation counts and coverage can differ between WoS and Scopus. Consequently, database-specific comparisons should be treated as coverage diagnostics rather than exact like-for-like quality comparisons.")
    sections=[
      ("4.3 Main Scientometric Indicators","overview",None,"Table 4.2 Main information about the collection"),
      ("4.4 Annual Scientific Production","annual","annual_production","Table 4.3 Annual publication and citation trend"),
      ("4.5 Most Productive Authors","authors","top_authors","Table 4.4 Most productive and influential authors"),
      ("4.6 Affiliation Analysis (College / University)","affiliations","top_affiliations","Table 4.5 Most productive colleges / universities"),
      ("4.7 Most Relevant Sources","sources","top_sources","Table 4.6 Most relevant publication sources"),
      ("4.8 Country Scientific Production","countries","top_countries","Table 4.7 Country-level scientific production"),
      ("4.9 Most Frequent Keywords","keywords","top_keywords","Table 4.8 Most frequent author and index keywords"),
      ("4.10 Document Types","document_types","document_types","Table 4.9 Distribution of document types"),
      ("4.11 Languages","languages","languages","Table 4.10 Distribution of publication languages"),
      ("4.12 Authorship and Collaboration","collaboration_summary",None,"Table 4.11 Authorship and collaboration indicators"),
      ("4.13 Bradford's Law","bradford",None,"Table 4.12 Bradford zones of source concentration"),
      ("4.14 Lotka's Law Distribution","lotka",None,"Table 4.13 Author productivity distribution"),
      ("4.15 Keyword Co-occurrence Network","keyword_nodes","keyword_network","Table 4.14 Keyword network centrality"),
      ("4.16 Co-authorship Network","author_nodes","author_network","Table 4.15 Co-authorship network centrality"),
      ("4.17 Topic Modelling","topics",None,"Table 4.16 Latent topics extracted from titles, abstracts and keywords"),
      ("4.18 WoS-Scopus Coverage Comparison","database_comparison",None,"Table 4.17 Database provenance comparison")]
    table_no=2
    for heading,key,figkey,caption in sections:
        doc.add_heading(heading,level=1)
        if key=="overview":
            d=results.get("main_information", pd.DataFrame())
            add_table(doc,d,caption,max_rows=100); doc.add_paragraph(interpretation("overview",overview=o))
        else:
            d=results.get(key,pd.DataFrame()); add_table(doc,d,caption)
            if figkey and figkey in figures:add_figure(doc,figures[figkey],f"Figure 4.{table_no-1} {heading.split(' ',1)[1]}")
            doc.add_paragraph(interpretation(key,d))
        table_no+=1
    doc.add_heading("4.19 Summary of Major Findings",level=1)
    doc.add_paragraph("The results collectively reveal the growth, intellectual structure, social structure, and citation influence of the research field. Publication and citation indicators identify scale and impact; author, source and country indicators identify leading contributors; collaboration networks identify structural relationships; keyword co-occurrence and topic modelling identify conceptual themes; and Bradford and Lotka distributions describe concentration and productivity patterns. These indicators should be interpreted as complementary evidence because no single scientometric measure captures research quality, novelty, societal value, and methodological rigor simultaneously.")
    doc.add_heading("4.20 Chapter Summary",level=1)
    doc.add_paragraph("This chapter transformed the combined WoS and Scopus metadata into an auditable set of descriptive, citation-based, network-based, and text-mining indicators. The generated tables and figures provide the empirical foundation for discussing dominant contributors, collaboration structures, influential publications, core journals, thematic clusters, and emerging topics in the succeeding chapter.")
    doc.save(outfile)
